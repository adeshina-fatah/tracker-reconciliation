"""Read shaded rows from the weekly workbook and write results back to a copy."""
from __future__ import annotations
from dataclasses import dataclass, field, asdict
from pathlib import Path
import shutil
import openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter
from .identifiers import clean, is_valid_job_no

@dataclass
class WorkItem:
    sheet: str
    row: int
    shade: str            # red | amber
    job_no: str | None
    rfq: str | None
    order_ref: str | None
    part_no: str | None
    customer: str | None
    status: str | None
    current_comment: str | None
    flags: list[str] = field(default_factory=list)

    @property
    def objective(self) -> str:
        return ("Find the customer purchase order (PO number, date received, value)."
                if self.shade == "amber" else
                "Find latest dispatch/ETA/BER-return evidence (status, dates, invoice, waybill).")

def _fill_rgb(cell) -> str | None:
    f = cell.fill
    if f is None or f.fill_type != "solid":
        return None
    c = f.start_color
    return c.rgb if c.type == "rgb" and isinstance(c.rgb, str) else None

def load_work_items(path: str | Path, cfg: dict) -> list[WorkItem]:
    wb = openpyxl.load_workbook(path, data_only=True)
    items: list[WorkItem] = []
    cols = cfg["identifier_columns"]
    for name in cfg["sheets"]:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        header = {str(c.value).strip(): i for i, c in enumerate(ws[1]) if c.value is not None}
        if cfg["comment_column"] not in header:
            continue
        ccol = header[cfg["comment_column"]]
        for row in ws.iter_rows(min_row=2):
            cell = row[ccol]
            shade = cfg["shades"].get(_fill_rgb(cell))
            if not shade:
                continue
            def col(key):
                i = header.get(cols[key]); return row[i].value if i is not None else None
            job_no = clean(col("job_no"))
            item = WorkItem(
                sheet=name, row=cell.row, shade=shade, job_no=job_no,
                rfq=clean(col("rfq")), order_ref=clean(col("order_ref")),
                part_no=clean(col("part_no")), customer=clean(col("customer")),
                status=clean(col("status")),
                current_comment=(str(cell.value).strip() or None) if cell.value else None,
            )
            if not is_valid_job_no(job_no):
                item.flags.append("job_no_format")
            if not job_no and not item.rfq and not item.order_ref:
                item.flags.append("no_identifiers")
            items.append(item)
    # duplicate job numbers
    seen = {}
    for it in items:
        if it.job_no:
            if it.job_no in seen:
                it.flags.append("duplicate_job_no"); seen[it.job_no].flags.append("duplicate_job_no")
            else:
                seen[it.job_no] = it
    return items

RESULT_COLUMNS = ["Proposed Comment", "Proposed Status", "Evidence", "Email Date",
                  "Source (Subject / Sender)", "Confidence", "Flags"]

def write_results(src: str | Path, dst: str | Path, results: list[dict], cfg: dict) -> Path:
    """Copy the workbook and append result columns beside Comments. Original untouched."""
    dst = Path(dst); dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(src, dst)
    wb = openpyxl.load_workbook(dst)
    by_sheet: dict[str, list[dict]] = {}
    for r in results:
        by_sheet.setdefault(r["sheet"], []).append(r)
    hdr_font = Font(bold=True)
    conf_fill = {"high": "FFC6EFCE", "medium": "FFFFEB9C", "low": "FFF4CCCC"}
    for sheet, rows in by_sheet.items():
        ws = wb[sheet]
        header = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
        start = header[cfg["comment_column"]] + 1
        for i, name in enumerate(RESULT_COLUMNS):
            c = ws.cell(row=1, column=start + i, value=name); c.font = hdr_font
            ws.column_dimensions[get_column_letter(start + i)].width = 28
        for r in rows:
            vals = [r.get("proposed_comment"), r.get("proposed_status"), r.get("evidence"),
                    r.get("email_date"), r.get("source"), r.get("confidence"), ", ".join(r.get("flags", []))]
            for i, v in enumerate(vals):
                cell = ws.cell(row=r["row"], column=start + i, value=v)
                if i == 5 and v in conf_fill:
                    cell.fill = PatternFill("solid", start_color=conf_fill[v])
    wb.save(dst)
    return dst

def items_to_records(items: list[WorkItem]) -> list[dict]:
    return [asdict(i) | {"objective": i.objective} for i in items]
